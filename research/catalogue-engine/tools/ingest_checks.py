"""Save a reviewer's spreadsheet checks so they survive the hourly workbook rebuild.

Usage: python3 tools/ingest_checks.py <checked.xlsx> <out.json> "<checked by>" <YYYY-MM-DD>
Reads the 'My check' (and optional 'Note') columns of the 'Research pins' sheet. Checks are reviewer
notes, not evidence: they never approve, verify or import a pin. Research Review records the
attributed human_direct_check when the reviewer confirms the pin there.
"""
import json, sys
from openpyxl import load_workbook

ALLOWED = {'ok', 'wrong', 'duplicate', 'not a pin', 'research needed'}


def read(path):
    ws = load_workbook(path, read_only=True)['Research pins']
    rows = ws.iter_rows(values_only=True)
    hdr = [str(h or '').strip().lower() for h in next(rows)]
    ci, ki = hdr.index('my check'), hdr.index('research_key')
    ni = hdr.index('note') if 'note' in hdr else None
    si = hdr.index('status') if 'status' in hdr else None
    out, bad = [], []
    for r in rows:
        v = str(r[ci] or '').strip()
        if not v: continue
        (out if v.lower() in ALLOWED else bad).append(dict(
            research_key=r[ki], check=v.lower(), note=(str(r[ni]).strip()[:240] if ni is not None and r[ni] else None),
            status_at_check=r[si] if si is not None else None))
    return out, bad


if __name__ == '__main__':
    src, dst, who, day = sys.argv[1:5]
    checks, bad = read(src)
    json.dump({'format': 'pinhunt-reviewer-checks-v1', 'checked_by': who, 'checked_on': day,
               'method': 'reviewer compared the listed source links with the record', 'checks': checks}, open(dst, 'w'), indent=1)
    print(f'{len(checks)} checks saved to {dst}; {len(bad)} rows with an unrecognised value')
    for b in bad[:20]: print('  unrecognised:', b['research_key'], b['check'])
